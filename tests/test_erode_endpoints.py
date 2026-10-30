import io
import openpyxl
from fastapi import status
from fastapi.testclient import TestClient
import pytest

from backend.app.main import app
from tests.fixtures.synthetic_erode import (
    create_synthetic_erode_bundle_pdf,
    create_synthetic_erode_transport_invoice_pdf,
    create_synthetic_erode_transport_pack_pdf,
)
from tests.fixtures.synthetic_tirunelveli import create_synthetic_tirunelveli_direct_pdf

client = TestClient(app, raise_server_exceptions=False)


def test_erode_extract_endpoint():
    """Verifies that the /erode/extract endpoint correctly extracts Erode documents."""
    pdf_bytes = create_synthetic_erode_bundle_pdf(is_direct=False)
    resp = client.post(
        "/api/v1/bookings/erode/extract?booking_type=group&enable_ocr=false",
        files={"file": ("erode_bundle.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    assert data["truck_number"] == "TN52L0517"
    assert data["loading_point"] == "erode"
    assert data["booking_type"] == "group"
    assert data["despatch_sheet"] is not None
    assert data["lr_sheet"] is not None
    assert len(data["despatch_sheet"]["rows"]) == 4
    assert len(data["lr_sheet"]["rows"]) == 3


def test_erode_reconcile_endpoint():
    """Verifies that the /erode/reconcile endpoint performs complete reconciliation."""
    pdf_bytes = create_synthetic_erode_bundle_pdf(is_direct=False)
    resp = client.post(
        "/api/v1/bookings/erode/reconcile?booking_type=group&enable_ocr=false",
        files={"file": ("erode_bundle.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    assert data["loading_point"] == "erode"
    assert len(data["truck_dossiers"]) == 1

    dossier = data["truck_dossiers"][0]
    assert len(dossier["reconciled_consignments"]) == 3
    assert dossier["truck_number"] == "TN52L0517"
    assert dossier["truck_total_packages"] == 60
    assert dossier["truck_total_weight_mt"] == 26.694


def test_erode_export_excel_endpoint():
    """Verifies end-to-end Excel generation for Erode bundles."""
    pdf_bytes = create_synthetic_erode_bundle_pdf(is_direct=False)
    resp = client.post(
        "/api/v1/bookings/erode/export-excel?booking_type=group&enable_ocr=false",
        files={"file": ("erode_bundle.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "TN52L0517" in resp.headers["content-disposition"]

    # Verify openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Consignments" in wb.sheetnames
    ws = wb["Consignments"]

    # Check banner
    banner = ws["A1"].value
    assert "CTLS TMS EXPORT - ERODE" in banner

    # Check row 3 (first consignment)
    truck_cell = ws["A3"].value
    lr_cell = ws["C3"].value
    inv_cell = ws["G3"].value
    wt_cell = ws["L3"].value
    frt_cell = ws["M3"].value

    assert truck_cell == "TN52L0517"
    assert lr_cell == "29738"
    assert inv_cell == "1250510477"
    assert wt_cell == 1.354
    assert frt_cell == 2505.0


def test_unified_export_endpoints():
    """Verifies that unified /export-excel routes to both Tirunelveli and Erode."""
    # 1. Unified with Erode
    erode_bytes = create_synthetic_erode_bundle_pdf(is_direct=True)
    resp1 = client.post(
        "/api/v1/bookings/export-excel?loading_point=erode&booking_type=direct&enable_ocr=false",
        files={"file": ("erode_direct.pdf", erode_bytes, "application/pdf")},
    )
    assert resp1.status_code == status.HTTP_200_OK
    assert "TN52AD9720" in resp1.headers["content-disposition"]

    # 2. Unified with Tirunelveli
    tiru_bytes = create_synthetic_tirunelveli_direct_pdf()
    resp2 = client.post(
        "/api/v1/bookings/export-excel?loading_point=tirunelveli&booking_type=direct&enable_ocr=false",
        files={"file": ("tiru_direct.pdf", tiru_bytes, "application/pdf")},
    )
    assert resp2.status_code == status.HTTP_200_OK
    assert "TN52Y7426" in resp2.headers["content-disposition"]


def test_erode_single_invoice_reconcile():
    """Verifies that an uploaded single Erode invoice produces confirmed consignments and page extractions."""
    pdf_bytes = create_synthetic_erode_transport_invoice_pdf()
    resp = client.post(
        "/api/v1/bookings/reconcile?loading_point=erode&booking_type=direct&enable_ocr=false",
        files={"file": ("erode_invoice_00481.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    assert data["loading_point"] == "erode"
    assert len(data["truck_dossiers"]) == 1
    dossier = data["truck_dossiers"][0]
    assert dossier["truck_number"] == "TN33AB1234"
    assert len(dossier["reconciled_consignments"]) == 1

    consignment = dossier["reconciled_consignments"][0]
    assert consignment["lr_number_raw"] == "ERD-LR-2026-019427"
    assert consignment["compressed_invoices_cell"] == "ERD/INV/2026/00481"
    assert consignment["total_value_rs"] == 56640.0

    # Ensure page extractions are present for UI inspect modal
    assert len(data["page_extractions"]) == 1
    p1 = data["page_extractions"][0]
    assert p1["classification"]["classification"] == "MATERIAL_INVOICE"
    assert p1["extracted_data"]["invoice_number"]["value"] == "ERD/INV/2026/00481"
    assert p1["extracted_data"]["printed_lr_number"]["raw_value"] == "ERD-LR-2026-019427"


def test_erode_transport_pack_reconcile_and_export():
    """Verifies that 4-page Erode combined test pack links LRs, Invoices, calculates freight and exports to Excel."""
    pdf_bytes = create_synthetic_erode_transport_pack_pdf()
    resp = client.post(
        "/api/v1/bookings/reconcile?loading_point=erode&booking_type=direct&enable_ocr=false",
        files={"file": ("erode_combined_test_pack.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    assert data["total_pages_in"] == 4
    assert data["total_pages_accounted_for"] == 4
    assert len(data["truck_dossiers"]) == 2

    # Verify dossier 1
    d1 = data["truck_dossiers"][0]
    assert d1["truck_number"] == "TN33AB1234"
    assert d1["truck_total_weight_mt"] == 1.24
    assert d1["truck_total_value_rs"] == 56640.0
    assert d1["truck_total_freight_rs"] == 7500.0
    c1 = d1["reconciled_consignments"][0]
    assert c1["lr_number_raw"] == "ERD-LR-2026-019427"
    assert c1["compressed_invoices_cell"] == "ERD/INV/2026/00481"
    assert c1["freight_amount"] == 7500.0

    # Verify dossier 2
    d2 = data["truck_dossiers"][1]
    assert d2["truck_number"] == "TN33CD5678"
    assert d2["truck_total_weight_mt"] == 0.68
    assert d2["truck_total_value_rs"] == 26550.0
    assert d2["truck_total_freight_rs"] == 3200.0
    c2 = d2["reconciled_consignments"][0]
    assert c2["lr_number_raw"] == "ERD-LR-2026-019428"
    assert c2["compressed_invoices_cell"] == "ERD/INV/2026/00482"
    assert c2["freight_amount"] == 3200.0

    # Verify Excel export
    export_resp = client.post(
        "/api/v1/bookings/export-excel?loading_point=erode&booking_type=direct&enable_ocr=false",
        files={"file": ("erode_combined_test_pack.pdf", pdf_bytes, "application/pdf")},
    )
    assert export_resp.status_code == status.HTTP_200_OK
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_erode_reconcile_without_booking_type_auto_detect():
    """Verifies that omitting booking_type automatically auto-classifies the booking per truck."""
    pdf_bytes = create_synthetic_erode_bundle_pdf(is_direct=False)
    # Notice: NO booking_type in query parameters!
    resp = client.post(
        "/api/v1/bookings/reconcile?loading_point=erode&enable_ocr=false",
        files={"file": ("erode_bundle.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    assert data["loading_point"] == "erode"
    assert len(data["truck_dossiers"]) == 1
    # Auto-detected as group booking because there are multiple LRs and multiple consignees
    assert data["truck_dossiers"][0]["booking_type"] == "group"
    assert len(data["truck_dossiers"][0]["reconciled_consignments"]) == 3



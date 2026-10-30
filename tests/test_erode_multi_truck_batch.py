import io
import os
from pathlib import Path
import openpyxl
import pytest
from fastapi import status
from fastapi.testclient import TestClient

from backend.app.main import app

client = TestClient(app, raise_server_exceptions=False)

PDF_PATH = Path("C:/Users/Admin/.gemini/antigravity/brain/b3befcc0-943f-4cfa-bc75-354343c8efd5/.user_uploaded/media_1793266363361_4def7bee.pdf")


@pytest.mark.skipif(not PDF_PATH.exists(), reason="24_01_26 Agreement.pdf artifact not found in test environment")
def test_erode_multi_truck_23_page_reconcile_and_export():
    """End-to-end verification of 23-page Erode multi-truck document bundle.

    Asserts:
    1. 23/23 pages strictly conserved (total_pages_in == 23, total_pages_accounted_for == 23).
    2. Exactly 10 discrete truck dossiers identified and segmented.
    3. Correct Direct vs Group booking classification across all trucks:
       - Trucks 1, 2, 4, 8, 9: GROUP
       - Trucks 3, 5, 6, 7, 10: DIRECT
    4. Over 25 confirmed consignments linked and reconciled.
    5. Formatted Excel export (.xlsx) generated with Trip Summary and Consignments sheets.
    """
    with open(PDF_PATH, "rb") as f:
        pdf_bytes = f.read()

    resp = client.post(
        "/api/v1/bookings/reconcile?loading_point=erode&booking_type=group&enable_ocr=true",
        files={"file": ("24_01_26_Agreement.pdf", pdf_bytes, "application/pdf")},
    )

    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()

    # 1. Page Conservation
    assert data["total_pages_in"] == 23
    assert data["total_pages_accounted_for"] == 23
    assert data["unassigned_pages"] == []

    # 2. 10 Truck Dossiers
    dossiers = data["truck_dossiers"]
    assert len(dossiers) == 10, f"Expected 10 truck dossiers, found {len(dossiers)}"

    # 3. Check individual truck classifications and consignments
    # Truck 1: TN34D5532 (P1, P2, P3) - GROUP (8 LRs)
    t1 = dossiers[0]
    assert "TN34D5532" in t1["truck_number"]
    assert t1["booking_type"] == "group"
    assert len(t1["reconciled_consignments"]) >= 6

    # Truck 2: TN81U1899 / TN21U1899 (P4, P5) - GROUP (4 LRs)
    t2 = dossiers[1]
    assert "1899" in t2["truck_number"]
    assert t2["booking_type"] == "group"
    assert len(t2["reconciled_consignments"]) >= 3

    # Truck 3: TN52K5837 (P6, P7) - DIRECT (1 LR)
    t3 = dossiers[2]
    assert "TN52K5837" in t3["truck_number"]
    assert t3["booking_type"] == "direct"
    assert len(t3["reconciled_consignments"]) == 1

    # Truck 4: TN59BA1080 (P8, P9) - GROUP (2 LRs)
    t4 = dossiers[3]
    assert "TN59BA1080" in t4["truck_number"]
    assert t4["booking_type"] == "group"
    assert len(t4["reconciled_consignments"]) >= 2

    # Truck 5: TN34R1832 (P10, P11) - DIRECT (1 LR)
    t5 = dossiers[4]
    assert "TN34R1832" in t5["truck_number"]
    assert t5["booking_type"] == "direct"
    assert len(t5["reconciled_consignments"]) == 1

    # Truck 6: TN52J9329 (P12, P13) - DIRECT (1 LR)
    t6 = dossiers[5]
    assert "TN52J9329" in t6["truck_number"]
    assert t6["booking_type"] == "direct"
    assert len(t6["reconciled_consignments"]) == 1

    # Truck 7: TN52AD9720 (P14, P15) - DIRECT (1 LR)
    t7 = dossiers[6]
    assert "TN52AD9720" in t7["truck_number"]
    assert t7["booking_type"] == "direct"
    assert len(t7["reconciled_consignments"]) == 1

    # Truck 8: TN52L0517 (P16, P17) - GROUP (3 LRs)
    t8 = dossiers[7]
    assert "TN52L0517" in t8["truck_number"]
    assert t8["booking_type"] == "group"
    assert len(t8["reconciled_consignments"]) >= 3

    # Truck 9: TN52H4793 (P18, P19, P20, P21) - GROUP (13 LRs)
    t9 = dossiers[8]
    assert "TN52H4793" in t9["truck_number"]
    assert t9["booking_type"] == "group"
    assert len(t9["reconciled_consignments"]) >= 10

    # Truck 10: TN52H7792 (P22, P23) - DIRECT (1 LR)
    t10 = dossiers[9]
    assert "TN52H7792" in t10["truck_number"]
    assert t10["booking_type"] == "direct"
    assert len(t10["reconciled_consignments"]) == 1

    total_consignments = sum(len(d["reconciled_consignments"]) for d in dossiers)
    assert total_consignments >= 25, f"Expected >= 25 total consignments, got {total_consignments}"

    # 4. Excel Generation Test
    export_resp = client.post(
        "/api/v1/bookings/export-excel-from-reconciliation",
        json=data,
    )
    assert export_resp.status_code == status.HTTP_200_OK
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "TN34D5532_batch_10trucks.xlsx" in export_resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(export_resp.content))
    assert "Consignments" in wb.sheetnames
    assert "Trip Summary & Gate Check" in wb.sheetnames
    assert "Page Provenance & Accounting" in wb.sheetnames

    # Check Trip Summary has 10 data rows (rows 3 to 12)
    summary_ws = wb["Trip Summary & Gate Check"]
    trucks_in_excel = [summary_ws.cell(row=r, column=1).value for r in range(3, 13)]
    assert len([t for t in trucks_in_excel if t]) == 10

